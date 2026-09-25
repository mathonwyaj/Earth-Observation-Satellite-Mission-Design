"""Build the Phase 10 preliminary design review report."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "final_design_review"
REPORT = OUT / "Earth_Observation_Satellite_Preliminary_Design_Review.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_border(cell, colour="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single"); tag.set(qn("w:sz"), size); tag.set(qn("w:color"), colour)
        borders.append(tag)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tc_mar.append(node)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.72); section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78); section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (("Title", 25, 0, 15), ("Heading 1", 17, 15, 7), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 4)):
        style = styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header = section.header.paragraphs[0]
    header.text = "EARTH OBSERVATION SATELLITE   |   PRELIMINARY DESIGN REVIEW"
    header.style = styles["Caption"]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.color.rgb = RGBColor(74, 85, 104); header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Mathonwy Akiwumi-Jones   |   Phase 10   |   ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(8); run.font.color.rgb = RGBColor(90, 90, 90)


def add_title_page(doc):
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.add_run("Earth Observation Satellite Mission and Systems Design")
    subtitle = doc.add_paragraph("Preliminary Design Review")
    subtitle.style = doc.styles["Subtitle"]
    subtitle.runs[0].font.size = Pt(17); subtitle.runs[0].font.color.rgb = RGBColor(45, 75, 105)
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2); meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [("Prepared by", "Mathonwy Akiwumi-Jones"), ("Review stage", "Phase 10 concept baseline"), ("Date", "25 September 2026"), ("Software", "Python 3, NumPy, Matplotlib and automated tests")]
    for row, (label, value) in zip(meta.rows, values):
        row.cells[0].text = label; row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells: set_cell_margins(cell, 80, 80, 80, 80); set_cell_border(cell, "FFFFFF", "0")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Purpose. ").bold = True
    p.add_run("This report records the selected preliminary design, the evidence generated across ten project phases, the final requirements status and the verification work that remains before any flight-design claim could be made.")
    doc.add_paragraph("Portfolio engineering study. Not flight-qualified hardware or guaranteed mission performance.").runs[0].italic = True
    doc.add_page_break()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = str(header); set_cell_shading(cell, "264B6A")
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(9)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for r_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value); cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_index % 2: set_cell_shading(cells[i], "F1F6FA")
            for run in cells[i].paragraphs[0].runs: run.font.size = Pt(9)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_border(cell); set_cell_margins(cell)
            if widths: cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, relative_path, caption, width=6.65):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / relative_path), width=Inches(width))
    c = doc.add_paragraph(caption); c.style = doc.styles["Caption"]; c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = False


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(item)


def build_report():
    doc = Document(); configure_document(doc); add_title_page(doc)

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph("The project defines a preliminary two-spacecraft multispectral Earth-observation mission for representative UK flood-monitoring use. Ten linked phases connect image requirements to payload sizing, radiometry, orbit and revisit, mission operations, spacecraft subsystem budgets, configuration, uncertainty propagation and final design review.")
    doc.add_paragraph("The final concept uses two spacecraft in a 550 km sun-synchronous orbit, a 10 m ground-sampling distance and an 81.92 km swath. A four-site polar ground-network concept reduces the modeled worst-case priority-delivery latency from 7.00 h to 1.39 h. Strengthened subsystem allocations raise simultaneous budget closure from 10.02% for the initial baseline to 98.48% in the reproducible 5,000-case Monte Carlo study.")
    doc.add_paragraph("The preliminary design review records 15 passed concept-level requirements, no failed requirements and one open item. Three-year mission life remains open because radiation, reliability, parts derating and lifetime testing are outside the current analytical model. The design is therefore accepted as a portfolio-level preliminary concept, not as a flight-qualified system.")
    add_table(doc, ["Review measure", "Result", "Status"], [
        ("Cumulative automated tests", "61 passed", "Verified"),
        ("Final simultaneous budget closure", "98.48%", "Pass against 95% target"),
        ("Worst priority-delivery latency", "1.39 h", "Pass against 3 h target"),
        ("Requirements verification", "15 pass, 1 open, 0 fail", "Conditional acceptance"),
    ], [2.5, 1.65, 2.3])

    doc.add_heading("Mission definition", level=1)
    doc.add_paragraph("The mission is an engineering reference case intended to demonstrate end-to-end space-systems reasoning. It is not tied to a customer contract. The top-level performance targets are spatial resolution, swath, revisit, delivery latency and a 60 kg wet-mass limit per spacecraft.")
    add_table(doc, ["Parameter", "Selected concept", "Design rationale"], [
        ("Architecture", "Two spacecraft", "Improves representative UK revisit"),
        ("Orbit", "550 km sun-synchronous", "Balances coverage, resolution and lighting"),
        ("Inclination", "97.593 deg", "Approximate J2 sun-synchronous solution"),
        ("Imaging", "10 m GSD, 81.92 km swath", "Meets mission targets"),
        ("Spectral bands", "Blue, green, red and NIR", "Flood and vegetation context"),
        ("Ground segment", "Four conceptual polar sites", "Closes preliminary 3 h delivery target"),
    ], [1.35, 2.15, 3.0])

    doc.add_heading("Engineering method", level=1)
    doc.add_paragraph("The analysis uses transparent first-order physical models linked through common design variables. Each phase produces machine-readable outputs and automated tests. The workflow is deliberately traceable: requirements drive sizing; sizing drives subsystem budgets; uncertainty analysis identifies weak margins; the final review strengthens allocations and rechecks closure.")
    add_table(doc, ["Phase", "Analysis", "Principal evidence"], [
        ("1 to 3", "Payload geometry, radiometry and optimisation", "Focal length, aperture, SNR, mass trade"),
        ("4", "Orbit and coverage", "SSO inclination, ground track, representative revisit"),
        ("5", "Operations and ground access", "Capacity, storage and latency"),
        ("6 to 7", "Subsystem sizing and integration", "Mass, power, link, ADCS, propulsion and thermal budgets"),
        ("8", "Spacecraft configuration", "Bus envelope, component layout, centre of mass and STL concept"),
        ("9", "Monte Carlo robustness", "5,000 cases and global sensitivity ranking"),
        ("10", "Final preliminary design review", "Selected allocations and final verification matrix"),
    ], [0.75, 2.45, 3.3])

    doc.add_heading("Payload design", level=1)
    doc.add_paragraph("An 8,192-pixel detector with 10 micrometre pitch gives a 550 mm focal length at 550 km altitude for 10 m nadir sampling. The resulting cross-track swath is 81.92 km. Radiometry and image-smear constraints were traded against aperture and exposure time. Phase 3 selected a 160 mm preliminary aperture; the final robustness allocation uses a 190 mm aperture proxy to protect worst-band SNR.")
    add_figure(doc, "results/optimisation/aperture_snr_optimisation.png", "Figure 1  Aperture and worst-band SNR optimisation")
    doc.add_paragraph("The aperture scaling remains a response-surface approximation. A detailed optical design must confirm modulation transfer function, stray light, focus stability, detector selection, baffle geometry and the mass and alignment effects of the larger telescope.")

    doc.add_heading("Orbit and coverage", level=1)
    doc.add_paragraph("The orbit model propagates circular Earth-fixed ground tracks and calculates the retrograde inclination required for approximate sun-synchronous nodal precession. A two-spacecraft architecture at 550 km with up to 25 deg off-nadir access produced a 37.79 h worst revisit among five representative UK sites, meeting the 48 h target.")
    add_figure(doc, "results/orbit_coverage/preferred_ground_track.png", "Figure 2  Preferred two-spacecraft ground-track representation")
    doc.add_paragraph("The revisit result applies to five sites rather than a complete UK-area coverage mask. Future mission analysis should use higher-fidelity propagation, cloud statistics, imaging constraints, exclusion zones and target-priority scheduling.")

    doc.add_heading("Mission operations and communications", level=1)
    doc.add_paragraph("Four representative 1,000 km scenes generate 39.32 Gbit per spacecraft per day after the assumed 4:1 compression. The original Goonhilly and Svalbard network had ample average capacity but a 7.00 h modeled worst delivery latency. Phase 10 adds conceptual Alaska and Antarctic polar sites; the idealised geometry gives 1.39 h worst latency and 1,163.09 Gbit/day mean capacity per spacecraft.")
    add_table(doc, ["Ground architecture", "Capacity per spacecraft", "Worst latency", "3 h target"], [
        ("Goonhilly", "143.37 Gbit/day", "9.94 h", "Fail"),
        ("Goonhilly and Svalbard", "562.28 Gbit/day", "7.00 h", "Fail"),
        ("Four-site polar concept", "1,163.09 Gbit/day", "1.39 h", "Pass"),
    ], [2.2, 1.75, 1.25, 1.1])
    doc.add_paragraph("This is an access-geometry result, not a contracted service assessment. It excludes simultaneous antenna conflicts, weather, maintenance, regulatory licensing, network routing, product processing time and commercial availability.")

    doc.add_heading("Spacecraft subsystem design", level=1)
    doc.add_paragraph("Phase 6 sized the initial power, battery, communications, attitude-control, propulsion, thermal and mass budgets. Phase 9 showed that several deterministic margins were too narrow under uncertainty. The final allocation therefore increases capability in the dominant risk areas.")
    add_table(doc, ["Subsystem", "Calculated baseline need", "Final allocation"], [
        ("Wet mass", "55.29 kg with system margin", "60 kg limit"),
        ("End-of-life solar power", "130.35 W", "190 W"),
        ("Battery", "53.47 Wh", "80 Wh"),
        ("Storage", "2.95 GB preliminary", "16 GB"),
        ("X-band link", "3.86 dB residual margin", "+4 dB design improvement"),
        ("Reaction-wheel torque", "0.00970 N m", "0.018 N m"),
        ("Reaction-wheel momentum", "0.0861 N m s", "0.180 N m s"),
        ("Mission delta-v", "48 m/s", "70 m/s"),
        ("Radiator area", "0.534 m2", "0.750 m2"),
    ], [1.75, 2.5, 2.25])
    add_figure(doc, "results/spacecraft_subsystems/subsystem_mass_budget.png", "Figure 3  Preliminary subsystem mass budget")

    doc.add_heading("Spacecraft configuration", level=1)
    doc.add_paragraph("The configuration model places the payload and major subsystems inside a 0.70 m by 0.70 m by 0.90 m bus. The deployed envelope is 2.20 m by 0.70 m by 0.90 m. The preliminary centre of mass lies approximately 19 mm from the bus centre. The configuration allocates 0.75 m2 of solar array and 0.60 m2 of radiator area in Phase 8; the final 0.75 m2 radiator allocation therefore requires a configuration update before detailed design.")
    add_figure(doc, "results/spacecraft_configuration/spacecraft_configuration_3d.png", "Figure 4  Preliminary spacecraft configuration")
    doc.add_paragraph("The generated STL is a packaging concept rather than a manufacturing CAD assembly. Structural load paths, mechanisms, harnessing, fasteners, keep-out zones, thermal interfaces and launch-vehicle accommodation remain outside its scope.")

    doc.add_heading("Robustness and sensitivity", level=1)
    doc.add_paragraph("The deterministic Phase 7 design passed every modeled subsystem budget but had little margin in several areas. Phase 9 propagated explicit uncertainty distributions across payload signal, mass, power, communications, attitude control, propulsion, thermal performance, data volume and ground availability. Only 10.02% of baseline cases passed every budget simultaneously.")
    add_figure(doc, "results/monte_carlo_robustness/robustness_and_sensitivity.png", "Figure 5  Baseline robustness and sensitivity evidence")
    doc.add_paragraph("The Phase 9 hardened comparison reached 92.66%. The final Phase 10 allocations increase aperture, power, communications, wheel, delta-v and radiator capability further, producing 98.48% simultaneous closure in the same 5,000 deterministic cases. Every individual modeled budget exceeds the 95% pass-rate target.")
    add_figure(doc, "results/final_design_review/final_design_review.png", "Figure 6  Final allocation robustness and requirements status")

    doc.add_heading("Requirements verification", level=1)
    doc.add_paragraph("The final matrix contains 16 tracked requirements. Fifteen pass through analysis, simulation or budget evidence. No requirement remains failed. Mission lifetime remains open because the current work does not contain the evidence needed to make a three-year reliability claim.")
    add_table(doc, ["ID", "Requirement", "Result", "Status"], [
        ("MIS-001", "10 m nadir sampling", "10.00 m", "Pass"),
        ("MIS-002", "At least 50 km swath", "81.92 km", "Pass"),
        ("PAY-001", "Worst-band SNR at least 50", "Final allocation passes", "Pass"),
        ("MIS-003", "UK-site revisit within 48 h", "37.79 h", "Pass"),
        ("MIS-004", "Priority delivery within 3 h", "1.39 h", "Pass"),
        ("SYS-001", "Wet mass no more than 60 kg", "55.29 kg baseline", "Pass"),
        ("ROB-001", "At least 95% robust closure", "98.48%", "Pass"),
        ("MIS-005", "Three-year mission life", "Analysis pending", "Open"),
    ], [0.85, 3.2, 1.55, 0.8])
    doc.add_paragraph("The complete matrix, including data, power, communications, ADCS, propulsion and thermal requirements, is supplied as final_requirements_matrix.csv in the Phase 10 results folder.")

    doc.add_heading("Verification evidence", level=1)
    doc.add_paragraph("The repository contains 61 cumulative automated tests. They check mathematical behavior, physical bounds, expected monotonic trends, reproducibility, output generation, budget closure and deliberate retention of open verification work. Passing tests demonstrate internal implementation consistency; they do not validate vendor hardware or replace environmental testing.")
    add_bullets(doc, [
        "Reproducible seed 260925 for the 5,000-case robustness study.",
        "Machine-readable CSV and JSON results for traceability.",
        "Plots generated directly from the analysis outputs.",
        "Explicit Pass, Fail and Open states rather than silently suppressing gaps.",
        "Cumulative unit-test discovery across all ten phases.",
    ])

    doc.add_heading("Limitations and remaining work", level=1)
    add_table(doc, ["Area", "Current limitation", "Required next evidence"], [
        ("Optics", "Response-surface aperture scaling", "Detailed optical prescription, MTF and tolerance model"),
        ("Coverage", "Five representative UK sites", "Area coverage, weather and scheduling simulation"),
        ("Ground segment", "Idealised station geometry", "Provider availability, conflicts, licensing and cost"),
        ("Structure", "Packaging STL only", "CAD assembly, launch-load FEA and modal verification"),
        ("Thermal", "Steady first-order sizing", "Hot and cold transient orbital model"),
        ("Radiation and reliability", "Not modeled", "Parts analysis, FMEA, derating and lifetime demonstration"),
        ("Operations", "No product-processing chain", "End-to-end scheduling and processing latency"),
        ("Cost and programme", "Outside present scope", "Cost estimate, schedule and procurement plan"),
    ], [1.45, 2.6, 2.55])

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("The project demonstrates a complete preliminary space-systems design loop: requirements were translated into payload and orbit choices, spacecraft and ground budgets were integrated, uncertainty exposed weak margins, and the final allocation was revised against measurable closure criteria. The selected concept meets all modeled performance and budget requirements, with 98.48% simultaneous Monte Carlo closure and 1.39 h modeled priority latency.")
    doc.add_paragraph("The correct review decision is conditional acceptance at preliminary concept level. The analytical design is coherent and reproducible, but three-year mission life and the detailed hardware, structural, optical, thermal and operational evidence remain future verification work. This distinction is central to the engineering value of the study.")

    doc.add_heading("Repository guide", level=1)
    add_table(doc, ["Command", "Purpose"], [
        ("python -m unittest discover -v", "Run all 61 cumulative tests"),
        ("python final_design_review.py", "Regenerate final review data and figure"),
        ("python build_final_report.py", "Regenerate this DOCX report"),
    ], [2.8, 3.8])
    if doc.paragraphs[-1].text == "":
        paragraph = doc.paragraphs[-1]._element
        paragraph.getparent().remove(paragraph)

    doc.core_properties.title = "Earth Observation Satellite Mission and Systems Design Preliminary Design Review"
    doc.core_properties.author = "Mathonwy Akiwumi-Jones"
    doc.core_properties.subject = "Phase 10 preliminary spacecraft systems design review"
    OUT.mkdir(parents=True, exist_ok=True); doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__": build_report()
